"""Verify one-page landscape summaries, text/links/bounds; MD remains the full report.
Requires pymupdf and python-pptx. Optional *.summary.json verifies every scene box.
"""
import argparse
import json
import re
from pathlib import Path
import fitz
from pptx import Presentation


def normalized(text):
    return re.sub(r"\s+", "", text)


def verify(directory, stem):
    markdown = (directory / f"{stem}.md").read_text(encoding="utf-8")
    pdf = fitz.open(directory / f"{stem}.pdf")
    deck = Presentation(directory / f"{stem}.pptx")
    assert len(pdf) == len(deck.slides) == 1, f"{stem}: must be exactly one page/slide"
    assert pdf[0].rect.width > pdf[0].rect.height
    assert abs(deck.slide_width / deck.slide_height - 16 / 9) < .001
    pdf_text = normalized(pdf[0].get_text())
    boxes = [s for s in deck.slides[0].shapes if s.has_text_frame]
    ppt_text = normalized(" ".join(s.text for s in boxes))
    for marker in ["표시", "수익률: 매수원금 대비 · 배당·수수료·세금 제외 · 억원/수익률 소수 둘째 자리 반올림 · 비주식 미수집", "총평가액", "평가손익", "총수익률"]:
        assert normalized(marker) in pdf_text and normalized(marker) in ppt_text, (stem, marker)
    footer = [s.text for s in boxes if s.top / 914400 >= 14.65]
    assert footer == ["수익률: 매수원금 대비 · 배당·수수료·세금 제외 · 억원/수익률 소수 둘째 자리 반올림 · 비주식 미수집"], (stem, footer)
    assert normalized("한 장 요약") not in pdf_text
    # Every displayed editable PPTX string must survive PDF font encoding and layout.
    for shape in boxes:
        assert normalized(shape.text) in pdf_text, (stem, "PDF missing", shape.text)
        rect = fitz.Rect(shape.left/12700, shape.top/12700, (shape.left+shape.width)/12700, (shape.top+shape.height)/12700)
        assert normalized(shape.text) in normalized(pdf[0].get_textbox(rect + (-1,-1,1,1))), (stem, "text box overflow", shape.text)
        if shape.top / 914400 >= 9 and shape.top / 914400 < 14.6 and not shape.text.startswith("게시"):
            # Names in overview rank columns remain full size, never autofit/shrink.
            if abs(shape.width / 914400 - 2.35) < .01:
                assert all(r.font.size.pt >= 16 for p in shape.text_frame.paragraphs for r in p.runs)
        for paragraph in shape.text_frame.paragraphs:
            for run in paragraph.runs:
                if run.font.size:
                    assert run.font.size.pt >= 10
    for page in pdf:
        for block in page.get_text("blocks"):
            assert 0 <= block[0] <= block[2] <= page.rect.width, (stem, block)
            assert 0 <= block[1] <= block[3] <= page.rect.height, (stem, block)
    for shape in deck.slides[0].shapes:
        assert 0 <= shape.left and shape.left + shape.width <= deck.slide_width
        assert 0 <= shape.top and shape.top + shape.height <= deck.slide_height
    scene_file = directory / f"{stem}.summary.json"
    if scene_file.exists():
        scene = json.loads(scene_file.read_text())
        for box in scene["texts"]:
            assert normalized(box["text"]) in pdf_text
            assert normalized(box["text"]) in ppt_text
            # Actual embedded-font metrics must stay inside assigned text boxes.
            rect = fitz.Rect(box["x"]*72, box["y"]*72, (box["x"]+box["w"])*72, (box["y"]+box["h"])*72)
            assert normalized(box["text"]) in normalized(pdf[0].get_textbox(rect + (-1,-1,1,1))), (stem, "box overflow", box)
    pdf_links = [link["uri"] for link in pdf[0].get_links() if "uri" in link]
    ppt_links = [r.hyperlink.address for s in boxes for p in s.text_frame.paragraphs for r in p.runs if r.hyperlink.address]
    assert set(pdf_links) == set(ppt_links)
    assert len(pdf_links) == sum(any(r.hyperlink.address for p in s.text_frame.paragraphs for r in p.runs) for s in boxes)
    assert all(url in markdown for url in pdf_links)
    pdf[0].get_pixmap(matrix=fitz.Matrix(1.5, 1.5)).save(directory/f"{stem}-page1.png")
    return {"stem": stem, "pages": len(pdf), "slides": len(deck.slides), "summaryTextBoxes": len(boxes), "linkedItems": len(pdf_links), "fullMarkdownRows": sum(line.startswith("- ") for line in markdown.splitlines()), "bounds": "pass"}


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("directory", type=Path)
    args = parser.parse_args()
    stems = [p.stem for p in sorted(args.directory.glob("*.pdf")) if (args.directory / (p.stem + ".pptx")).exists()]
    assert stems, "No PDF/PPTX pairs found"
    result = [verify(args.directory, stem) for stem in stems]
    (args.directory/"export-verification.json").write_text(json.dumps(result, indent=2), encoding="utf-8")
    print(json.dumps(result, indent=2))
